"""Test penyusunan laporan Excel & ekspor CSV."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from valido.engine import run_audit
from valido.report import (
    build_report,
    export_clean_csv,
    export_findings_csv,
)


@pytest.fixture
def result(messy_df, sample_config):
    return run_audit(messy_df, sample_config)


def test_build_report_creates_all_sheets(result, tmp_path):
    path = build_report(result, tmp_path / "laporan.xlsx", notes=["catatan uji"])
    assert path.exists() and path.stat().st_size > 5000

    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Ringkasan",
        "Temuan_Detail",
        "Statistik_Aturan",
        "Profil_Kolom",
        "Data_Bersih",
        "Konfigurasi",
    ]


def test_summary_sheet_contents(result, tmp_path):
    path = build_report(result, tmp_path / "laporan.xlsx")
    ws = load_workbook(path)["Ringkasan"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)

    assert "LAPORAN AUDIT KUALITAS DATA" in text
    assert f"{result.summary['total_rows']}" in text
    assert "SKOR KESEHATAN DATA" in text
    assert "TEMUAN PER TINGKAT KEPARAHAN" in text
    assert "SATU" not in text  # placeholder guard


def test_findings_sheet_has_rows(result, tmp_path):
    path = build_report(result, tmp_path / "laporan.xlsx")
    ws = load_workbook(path)["Temuan_Detail"]
    header_cells = [c.value for c in ws[3] if c.value]
    assert "Saran Perbaikan" in header_cells
    assert "Alasan" in header_cells
    assert ws.max_row >= len(result.findings)


def test_report_without_findings(clean_df, clean_config, tmp_path):
    result = run_audit(clean_df, clean_config)
    path = build_report(result, tmp_path / "bersih.xlsx")
    wb = load_workbook(path)
    assert "Temuan_Detail" in wb.sheetnames
    ws = wb["Temuan_Detail"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Tidak ada temuan" in text


def test_report_respects_clean_preview_toggle(result, tmp_path):
    path = build_report(result, tmp_path / "tanpa_bersih.xlsx", include_clean_preview=False)
    wb = load_workbook(path)
    assert "Data_Bersih" not in wb.sheetnames
    assert "Konfigurasi" in wb.sheetnames


def test_config_sheet_contains_yaml(result, tmp_path):
    path = build_report(result, tmp_path / "laporan.xlsx")
    ws = load_workbook(path)["Konfigurasi"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "rules:" in text
    assert "wajib_nama" in text


def test_export_findings_csv(result, tmp_path):
    path = export_findings_csv(result, tmp_path / "temuan.csv")
    df = pd.read_csv(path, encoding="utf-8-sig")
    assert len(df) == len(result.findings)
    assert "Saran Perbaikan" in df.columns


def test_export_clean_csv(result, tmp_path):
    path = export_clean_csv(result, tmp_path / "bersih.csv")
    df = pd.read_csv(path, encoding="utf-8-sig")
    assert df.shape[0] == len(result.df)


def test_report_creates_parent_directories(result, tmp_path):
    target = tmp_path / "sub" / "folder" / "laporan.xlsx"
    path = build_report(result, target)
    assert path.exists()


def test_report_metadata(result, tmp_path):
    path = build_report(result, tmp_path / "laporan.xlsx")
    wb = load_workbook(path)
    assert wb.properties.creator
    assert "Valido" in (wb.properties.description or "") or "audit" in (wb.properties.description or "").lower()
